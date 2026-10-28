#!/usr/bin/env python3
"""Generate the AR/VR guided assembly seminar deck, mirroring the senior's layout."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

OUT = "/home/user/notebooklm/reports/ARVR_Guided_Assembly_Seminar_Deck.pptx"

# theme
NAVY = RGBColor(0x14, 0x2B, 0x45)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
ACCENT = RGBColor(0xF2, 0xA7, 0x2E)
LIGHT = RGBColor(0xF4, 0xF6, 0xF8)
GREY = RGBColor(0x5B, 0x6B, 0x7B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1C, 0x24, 0x2E)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, color, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=6, line_spacing=1.08, wrap=True):
    """runs: list of paragraphs; each paragraph is list of (text,size,bold,color,italic)."""
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, pruns in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (t, sz, b, col, it) in pruns:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.bold = b; r.font.italic = it
            r.font.color.rgb = col; r.font.name = "Calibri"
    return tb


def bullets(s, x, y, w, h, items, size=16, color=DARK, gap=7, lh=1.1):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.space_before = Pt(0); p.line_spacing = lh
        # bullet glyph
        r0 = p.add_run(); r0.text = "•  "
        r0.font.size = Pt(size); r0.font.bold = True; r0.font.color.rgb = TEAL; r0.font.name = "Calibri"
        if isinstance(it, tuple):
            lead, rest = it
            r1 = p.add_run(); r1.text = lead
            r1.font.size = Pt(size); r1.font.bold = True; r1.font.color.rgb = NAVY; r1.font.name = "Calibri"
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(size); r2.font.color.rgb = color; r2.font.name = "Calibri"
        else:
            r1 = p.add_run(); r1.text = it
            r1.font.size = Pt(size); r1.font.color.rgb = color; r1.font.name = "Calibri"
    return tb


def content_header(s, title, kicker=None):
    rect(s, 0, 0, SW, Inches(1.25), NAVY)
    rect(s, 0, Inches(1.25), SW, Pt(4), ACCENT)
    txt(s, Inches(0.6), Inches(0.28), Inches(11.5), Inches(0.75),
        [[(title, 28, True, WHITE, False)]], anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        txt(s, Inches(10.2), Inches(0.32), Inches(2.6), Inches(0.6),
            [[(kicker, 13, True, ACCENT, False)]], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def citation(s, text):
    txt(s, Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.4),
        [[("Source: " + text, 10.5, False, GREY, True)]], anchor=MSO_ANCHOR.MIDDLE)


def footer_tag(s, n):
    txt(s, Inches(11.9), Inches(7.02), Inches(1.2), Inches(0.35),
        [[(f"{n:02d}", 11, True, GREY, False)]], align=PP_ALIGN.RIGHT)


# ================= SLIDE 1 : TITLE =================
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(2.1), SW, Pt(3), ACCENT)
rect(s, 0, Inches(5.05), SW, Pt(3), TEAL)
txt(s, Inches(0.9), Inches(0.55), Inches(11.5), Inches(0.5),
    [[("SEMINAR & TECHNICAL REVIEW", 16, True, ACCENT, False)]])
txt(s, Inches(0.9), Inches(2.35), Inches(11.6), Inches(2.4),
    [[("An AR/VR and AI-Guided Manual", 40, True, WHITE, False)],
     [("Assembly Solution for India", 40, True, WHITE, False)]],
    line_spacing=1.05, space_after=4)
txt(s, Inches(0.9), Inches(5.35), Inches(11.6), Inches(1.6),
    [[("By,", 17, False, LIGHT, False)],
     [("« Your Name »,  Roll No. « 24MEM3Rxx »,", 18, True, WHITE, False)],
     [("M.Tech CIM  —  Department of Mechanical Engineering", 15, False, LIGHT, False)],
     [("National Institute of Technology Warangal  ·  2025-26", 14, False, RGBColor(0xB9,0xC6,0xD4), False)]],
    space_after=5)

# ================= SLIDE 2 : TABLE OF CONTENTS =================
s = slide()
rect(s, 0, 0, SW, SH, WHITE)
rect(s, 0, 0, Inches(4.5), SH, NAVY)
txt(s, Inches(0.55), Inches(0.9), Inches(3.5), Inches(2.0),
    [[("TABLE", 40, True, WHITE, False)], [("OF", 40, True, ACCENT, False)], [("CONTENTS", 40, True, WHITE, False)]],
    line_spacing=1.0, space_after=2)
toc = [
    ("01", "Introduction"),
    ("02", "Problem & Opportunity"),
    ("03", "Patent Portfolio (5 patents)"),
    ("04", "Unique Value Proposition"),
    ("05", "Market & Business Model"),
    ("06", "Financials, Roadmap & Conclusion"),
]
y = Inches(0.95)
for num, label in toc:
    txt(s, Inches(5.0), y, Inches(1.0), Inches(0.8),
        [[(num, 30, True, TEAL, False)]], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(6.1), y, Inches(6.7), Inches(0.8),
        [[(label, 20, True, NAVY, False)]], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.98)

# ================= SLIDE 3 : INTRODUCTION =================
s = slide(); content_header(s, "Introduction", "01")
bullets(s, Inches(0.6), Inches(1.65), Inches(12.1), Inches(5.0), [
    ("Manual assembly dominates Indian manufacturing — ", "automotive harnesses, switchgear, pumps and electronics are built by hand across many low-volume variants."),
    ("Recurring pain: ", "wrong or missed parts, wrong sequence, slow onboarding, out-of-date paper SOPs, and no per-unit traceability."),
    ("AR / VR / computer vision help: ", "AR shows the next step on the real work, VR trains before touching parts, and AI checks each step from a camera."),
    ("Proven effect: ", "overlaying 3D instructions on the work cut assembly errors by 82% in a controlled study (Tang et al., CHI 2003)."),
    ("Patent-informed, not patent-protected: ", "the architecture draws on five real patents' public disclosures, kept separate from their narrower granted claims."),
], size=17, gap=11)
citation(s, "A. Tang et al., “Comparative effectiveness of AR in object assembly,” CHI 2003, pp. 73-80.")
footer_tag(s, 3)

# ================= SLIDE 4 : PROBLEM & OPPORTUNITY =================
s = slide(); content_header(s, "Problem & Opportunity", "02")
# two columns
txt(s, Inches(0.6), Inches(1.5), Inches(6.0), Inches(0.5), [[("The problem", 18, True, TEAL, False)]])
bullets(s, Inches(0.6), Inches(2.05), Inches(6.1), Inches(4.5), [
    ("Human error: ", "wrong part, missed screw, wrong torque → rework, scrap, OEM PPM penalties."),
    ("Onboarding & attrition: ", "shop-floor attrition >18-20%; 80% of employers report a skills shortage."),
    ("Stale paper SOPs: ", "rarely read during work; many variants confuse the station."),
    ("Headset / platform risk: ", "HoloLens 2 ended 2024; Dynamics 365 Guides ends 31 Dec 2026."),
], size=15.5, gap=8)
txt(s, Inches(6.95), Inches(1.5), Inches(5.8), Inches(0.5), [[("The opportunity", 18, True, TEAL, False)]])
bullets(s, Inches(6.95), Inches(2.05), Inches(5.8), Inches(4.5), [
    ("Auto components: ", "₹6.73 lakh crore turnover FY25 (ACMA), 14% 5-yr CAGR."),
    ("Electronics: ", "~₹11.3 lakh crore production FY25 (Economic Survey)."),
    ("Migration gap: ", "Guides retirement leaves customers seeking a replacement."),
    ("Policy tailwind: ", "Make in India, Industry 4.0, Startup India Seed Fund (≤₹20 L)."),
], size=15.5, gap=8)
citation(s, "ACMA FY25 Industry Performance Review; Microsoft Learn (Guides retirement).")
footer_tag(s, 4)

# ================= SLIDES 5-9 : PATENTS =================
patent_slides = [
    ("Patent 1: Projected-Light Guided Assembly", "US 7,515,981 B2  ·  Granted  ·  LightGuide, Inc.",
     "Inventors: Paul Ryznar, James Ryznar  ·  Priority 2005  ·  US term ~2026 (near expiry); later family runs to ~2032",
     ["Projects arrows, highlighted bins, locations and text directly onto the work area.",
      "Steps the operator through the assembly in order, each step confirmed before the next.",
      "Headset-free, station-mounted — the model StepSure follows for the shop floor.",
      "Expiry of the 2005 patent does NOT free the idea: the later family is in force and litigated."],
     "US 7,515,981 B2, “Light guided assembly system,” patents.google.com/patent/US7515981B2."),
    ("Patent 2: AR Assembly Assistance & Monitoring", "US 10,366,521 B1  ·  Granted  ·  Amazon Technologies, Inc.",
     "Inventors: J. Peacock, J. Watson, S. Wilcox  ·  Filed ~2017  ·  Expected term to ~2037",
     ["AR glasses show the information for the current step to the worker.",
      "Camera image data verifies the assembly step; safe paths and restricted zones are overlaid.",
      "Shows guidance + verification belong together — core to StepSure's design.",
      "Prior art: AR guidance + verification is known, so novelty must be in HOW it is verified."],
     "US 10,366,521 B1, “Augmented reality assembly assistance and monitoring,” Amazon, 2019."),
    ("Patent 3: ML Assembly Step Recognition", "US 11,715,300 B1  ·  Granted  ·  assignee likely Robert Bosch GmbH",
     "Inventors: Y. Guo, X. Huang, L. Ren  ·  Filed 2022  ·  Expected term to ~2042",
     ["ML model trained on data from 2D images + 3D CAD turns camera frames into feature vectors.",
      "Template matching finds base and incoming parts and the current step; a 2nd model verifies it.",
      "An AR position indicator shows where the part goes, updating as new frames arrive.",
      "Closest to StepSure's AI core — and its BIGGEST freedom-to-operate risk; needs a different method or a licence."],
     "US 11,715,300 B1, “Product assembly step recognition using AR,” 2023."),
    ("Patent 4: VR Assembly Training Simulator", "US 12,592,164 B2  ·  Granted  ·  Permco, Inc.",
     "Filed 2022  ·  Granted 2026  ·  Expected term to ~2042",
     ["VR headset presents a virtual workbench, tools and components.",
      "Shows an ordered list of steps and highlights the tool and location for each.",
      "A haptic glove teaches workers to feel surface defects such as burrs.",
      "Matches StepSure's VR-training half — one recipe authored once drives both training and floor guidance."],
     "US 12,592,164 B2, “Virtual reality training simulator,” Permco, 2026."),
    ("Patent 5: AR/VR Interaction with Cobots", "US 11,529,737 B2  ·  Granted  ·  Raytheon Company",
     "Inventor: Kristen M. Stone  ·  EP family 4096874  ·  Priority 2020, granted 2022",
     ["A wearable AR/VR headset shows the cobot's planned motion path before it moves.",
      "A VR training simulation feeds information to the workstation to help program the path.",
      "Covers StepSure's future human-cobot extension (robot-intent overlay, hand-over).",
      "The overlay is an INFORMATION aid, not a safety function (ISO 10218 / ISO TS 15066 still apply)."],
     "US 11,529,737 B2, “VR/AR interaction with collaborative robots,” Raytheon, 2022."),
]
for idx, (title, biblio, meta, pts, cite) in enumerate(patent_slides, start=5):
    s = slide(); content_header(s, title, "03")
    # biblio band
    rect(s, Inches(0.6), Inches(1.55), Inches(12.1), Inches(0.95), LIGHT)
    rect(s, Inches(0.6), Inches(1.55), Pt(5), Inches(0.95), TEAL)
    txt(s, Inches(0.85), Inches(1.64), Inches(11.7), Inches(0.9),
        [[(biblio, 15.5, True, NAVY, False)],
         [(meta, 12.5, False, GREY, True)]], space_after=3, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, Inches(0.75), Inches(2.85), Inches(11.9), Inches(3.8), pts, size=16, gap=10)
    citation(s, cite)
    footer_tag(s, idx)

# ================= SLIDE 10 : UVP =================
s = slide(); content_header(s, "Unique Value Proposition", "04")
cards = [
    ("Checked steps", "Camera + tools verify every step before the next unlocks — not just shown."),
    ("Headset-free", "Fixed camera + projector/screen; shared across shifts; no vendor lock-in."),
    ("One recipe, 3 outputs", "One process drives guidance, CV checking and auto-built VR training."),
    ("Built for India", "Icon-first UI, Indian-language voice, INR pricing, cluster-based support."),
    ("Cobot-ready", "Same recipe extends to a future robot-intent overlay and hand-over."),
]
x0, y0, cw, chh, gapx = Inches(0.6), Inches(1.7), Inches(3.86), Inches(2.3), Inches(0.26)
positions = [(0,0),(1,0),(2,0),(0,1),(1,1)]
for (col,row),(t,d) in zip(positions, cards):
    x = x0 + col*(cw+gapx); y = y0 + row*(chh+Inches(0.3))
    rect(s, x, y, cw, chh, LIGHT)
    rect(s, x, y, cw, Pt(6), ACCENT)
    txt(s, x+Inches(0.22), y+Inches(0.28), cw-Inches(0.44), Inches(0.6),
        [[(t, 18, True, NAVY, False)]])
    txt(s, x+Inches(0.22), y+Inches(0.95), cw-Inches(0.44), Inches(1.2),
        [[(d, 13.5, False, DARK, False)]], line_spacing=1.1)
txt(s, Inches(4.75), Inches(6.55), Inches(8.0), Inches(0.6),
    [[("“Every step checked, every worker productive in days — without headsets, at an Indian price.”", 14.5, True, TEAL, True)]],
    anchor=MSO_ANCHOR.MIDDLE)
footer_tag(s, 10)

# ================= SLIDE 11 : MARKET ANALYSIS =================
s = slide(); content_header(s, "Market Analysis", "05")
# stat tiles
stats = [("₹6.73 L cr", "Auto-component turnover, FY25 (ACMA)"),
         ("~₹11.3 L cr", "Electronics production, FY25"),
         ("~400", "Target ICP accounts, 4 clusters")]
x = Inches(0.6)
for val, lab in stats:
    rect(s, x, Inches(1.6), Inches(3.9), Inches(1.35), NAVY)
    txt(s, x+Inches(0.2), Inches(1.72), Inches(3.5), Inches(0.7), [[(val,26,True,ACCENT,False)]])
    txt(s, x+Inches(0.2), Inches(2.42), Inches(3.5), Inches(0.5), [[(lab,12.5,False,WHITE,False)]], line_spacing=1.0)
    x += Inches(4.07)
txt(s, Inches(0.6), Inches(3.2), Inches(6.0), Inches(0.5), [[("Competitive landscape", 17, True, TEAL, False)]])
bullets(s, Inches(0.6), Inches(3.7), Inches(6.2), Inches(3.2), [
    ("LightGuide: ", "mature projected AR; India pricing unverified."),
    ("Retrocausal: ", "AI action recognition; guidance/VR not core."),
    ("PTC Vuforia: ", "enterprise AR authoring, per-user quote."),
    ("AutoVRse (India): ", "VR training-led, not live line checking."),
], size=14.5, gap=7)
txt(s, Inches(7.0), Inches(3.2), Inches(5.8), Inches(0.5), [[("Our position & segments", 17, True, TEAL, False)]])
bullets(s, Inches(7.0), Inches(3.7), Inches(5.8), Inches(3.2), [
    "Verified guidance + VR from one recipe, headset-free, INR.",
    "Tier-1/2 auto suppliers (pilot segment).",
    "Electronics / EMS and electrical.",
    "OEM plants; ITIs & training centres (VR seats).",
], size=14.5, gap=7)
citation(s, "ACMA FY25; Knight Frank / industry reports; company disclosures.")
footer_tag(s, 11)

# ================= SLIDE 12 : BUSINESS MODEL =================
s = slide(); content_header(s, "Business Model", "05")
txt(s, Inches(0.6), Inches(1.5), Inches(12.0), Inches(0.5),
    [[("Hybrid: per-station subscription (main) + hardware near cost + onboarding + VR seats. All prices are assumptions.", 14.5, False, GREY, True)]])
streams = [
    ("Station subscription", "₹20,000 / station / month", "Main recurring revenue (RaaS-style)"),
    ("Hardware kit", "~₹1.6 lakh one-time", "Sold near ₹1.35 L cost to cut barrier"),
    ("Onboarding / recipe", "~₹3 lakh per customer", "Process digitisation + MES/ERP link"),
    ("VR training licence", "₹30,000 / headset / year", "For ITIs & training departments"),
    ("Paid pilot (entry)", "~₹7 lakh / 2 stations / 12 wks", "Screens prospects, funds case study"),
    ("Cobot Bridge (future)", "~₹5,000 / station / month", "After FTO clearance"),
]
x0, y0, cw, chh, gx, gy = Inches(0.6), Inches(2.25), Inches(3.9), Inches(1.75), Inches(0.26), Inches(0.3)
for i,(t,price,note) in enumerate(streams):
    col, row = i%3, i//3
    x = x0+col*(cw+gx); y=y0+row*(chh+gy)
    rect(s, x, y, cw, chh, LIGHT)
    rect(s, x, y, Pt(5), chh, TEAL)
    txt(s, x+Inches(0.2), y+Inches(0.16), cw-Inches(0.4), Inches(0.5), [[(t,15.5,True,NAVY,False)]])
    txt(s, x+Inches(0.2), y+Inches(0.64), cw-Inches(0.4), Inches(0.5), [[(price,16,True,ACCENT,False)]])
    txt(s, x+Inches(0.2), y+Inches(1.14), cw-Inches(0.4), Inches(0.5), [[(note,12,False,GREY,False)]], line_spacing=1.0)
footer_tag(s, 12)

# ================= SLIDE 13 : FINANCIALS =================
s = slide(); content_header(s, "Financial Projections", "06")
txt(s, Inches(0.6), Inches(1.45), Inches(12.0), Inches(0.4),
    [[("Every figure is a labelled assumption, not a forecast. Standard customer = 5 stations.", 13.5, False, GREY, True)]])
# table
rows = [
    ["Metric", "Year 1 (Dev & Pilot)", "Year 2 (Scale-Up)", "Year 3 (Break-Even)"],
    ["Revenue", "₹0.2-0.3 Cr", "₹1.2-1.5 Cr", "₹3-4 Cr"],
    ["Gross margin", "~45%", "~65%", "~70%"],
    ["Key driver", "2 paid pilots", "3-5 customers", "10-15 customers"],
    ["Break-even", "N/A", "N/A", "~12-13 customers"],
]
tx, ty, tw = Inches(0.6), Inches(2.0), Inches(12.1)
rowh = Inches(0.62); cw0 = Inches(2.4); cwo = (tw-cw0)/3
for ri, row in enumerate(rows):
    y = ty + ri*rowh
    for ci, val in enumerate(row):
        x = tx + (0 if ci==0 else cw0 + (ci-1)*cwo)
        w = cw0 if ci==0 else cwo
        if ri==0:
            rect(s, x, y, w, rowh, NAVY)
            col = WHITE; b=True; sz=14
        else:
            rect(s, x, y, w, rowh, LIGHT if ri%2 else WHITE, line=RGBColor(0xD8,0xDE,0xE4))
            col = NAVY if ci==0 else DARK; b = (ci==0) or (ri==4 and ci==3); sz=13.5
            if ri==4 and ci==3: col = TEAL
        txt(s, x+Inches(0.12), y, w-Inches(0.24), rowh, [[(val,sz,b,col,False)]], anchor=MSO_ANCHOR.MIDDLE)
bullets(s, Inches(0.6), Inches(5.3), Inches(12.1), Inches(1.6), [
    ("Year-1 funding need ≈ ₹1.1 crore; ", "break-even ≈ 1.2 Cr fixed ÷ 9.6 L gross profit per customer ≈ 12.5 customers (~62 stations)."),
    ("Most sensitive lever: ", "stations per customer — at 10/customer break-even drops to ~6-7. Not profitable in year 1 in any scenario."),
], size=13.5, gap=6)
footer_tag(s, 13)

# ================= SLIDE 14 : ROADMAP & CONCLUSION =================
s = slide(); content_header(s, "Implementation Roadmap & Conclusion", "06")
phases = [
    ("Phase 1", "0-6 mo", "Prototype: 1 station, 3 check types, ≥95% accuracy; VR trainer; seed funding."),
    ("Phase 2", "6-12 mo", "Paid pilot at 1 plant; measure KPIs vs baseline; convert + case study."),
    ("Phase 3", "12-24 mo", "v1.0: Recipe Studio, dashboard, MES export, multilingual; customers 2-3."),
    ("Phase 4", "24-36 mo", "Cobot Bridge (post-FTO); break-even; indigenous patents; SE Asia."),
]
x0 = Inches(0.6); cw = Inches(3.0); gx = Inches(0.17); y = Inches(1.7)
for i,(ph,tm,desc) in enumerate(phases):
    x = x0 + i*(cw+gx)
    rect(s, x, y, cw, Inches(0.55), TEAL)
    txt(s, x+Inches(0.15), y, cw-Inches(0.3), Inches(0.55),
        [[(ph+"  ", 16, True, WHITE, False),(tm, 13, False, RGBColor(0xCF,0xEA,0xEC), False)]], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, x, y+Inches(0.55), cw, Inches(2.15), LIGHT)
    txt(s, x+Inches(0.18), y+Inches(0.72), cw-Inches(0.36), Inches(1.9),
        [[(desc, 13.5, False, DARK, False)]], line_spacing=1.12)
rect(s, Inches(0.6), Inches(4.85), Inches(12.1), Inches(1.9), NAVY)
txt(s, Inches(0.9), Inches(5.0), Inches(11.5), Inches(0.5), [[("Conclusion", 18, True, ACCENT, False)]])
txt(s, Inches(0.9), Inches(5.5), Inches(11.5), Inches(1.2),
    [[("The gap is not AR for assembly — it already exists — but an affordable, headset-free, "
       "check-every-step system reusing one recipe for training, guidance and future cobot work. StepSure "
       "tests it with one paid pilot, India-first. The decisive unknowns are willingness to pay and real-line "
       "CV accuracy; a professional freedom-to-operate opinion is required before any build or sale.",
       13.5, False, WHITE, False)]], line_spacing=1.15)
footer_tag(s, 14)

# ================= SLIDE 15 : THANK YOU =================
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, Inches(0.9), Inches(3.55), Inches(2.2), Pt(5), ACCENT)
txt(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.2), [[("Thank You!", 54, True, WHITE, False)]])
txt(s, Inches(0.9), Inches(3.8), Inches(11.5), Inches(0.8),
    [[("Questions & discussion welcome.", 18, False, LIGHT, False)]])
txt(s, Inches(0.9), Inches(6.6), Inches(11.5), Inches(0.5),
    [[("« Your Name »  ·  M.Tech CIM  ·  NIT Warangal  ·  2025-26", 13, False, RGBColor(0xB9,0xC6,0xD4), False)]])

prs.save(OUT)
print("saved", OUT, "slides", len(prs.slides._sldIdLst))
